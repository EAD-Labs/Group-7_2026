"""
Pi Jam Foundation Classroom Assessment System — Export Service
Generates styled Excel workbooks and merged PDF dossiers.
"""

import os
from typing import Optional, Dict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from pypdf import PdfWriter

from backend.database import get_db, EXPORT_DIR
from backend.services.analytics_service import calculate_session_analysis


def generate_excel_scorecard(session_id: str, custom_answer_key: Optional[Dict[str, str]] = None) -> str:
    """
    Generates a production-styled Excel spreadsheet (.xlsx) for the assessment session.
    Contains 3 formatted sheets: Executive Summary, Student Scorecards, and Item Analysis.
    Returns the absolute path to the generated Excel file.
    """
    data = calculate_session_analysis(session_id, custom_answer_key)
    wb = openpyxl.Workbook()

    # Brand Colors
    PRIMARY_COLOR = "1A365D"  # Deep Navy
    SECONDARY_COLOR = "2B6CB0" # Slate Blue
    CORRECT_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid") # Soft Green
    WRONG_FILL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")   # Soft Red
    HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color=SECONDARY_COLOR, end_color=SECONDARY_COLOR, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    # -------------------------------------------------------------
    # SHEET 1: Executive Summary
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary.merge_cells("A1:E1")
    ws_summary["A1"] = "Pi Jam Foundation — Formative Assessment Summary"
    ws_summary["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws_summary["A1"].fill = HEADER_FILL
    ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 35

    # Metadata Details
    meta = [
        ("Quiz Name:", data["quiz_name"]),
        ("Subject:", data["subject"]),
        ("Class & Section:", data["class_label"]),
        ("School:", data["school_name"]),
        ("Teacher:", data["teacher_name"]),
        ("Date Conducted:", data["session_date"]),
        ("Question Count:", f"{data['question_count']} Questions"),
        ("Question Paper Mode:", data["question_mode"].upper()),
    ]

    for idx, (label, val) in enumerate(meta, start=3):
        ws_summary[f"B{idx}"] = label
        ws_summary[f"B{idx}"].font = Font(bold=True, color="2D3748")
        ws_summary[f"C{idx}"] = val
        ws_summary[f"C{idx}"].font = Font(color="1A202C")
        ws_summary[f"B{idx}"].border = thin_border
        ws_summary[f"C{idx}"].border = thin_border

    # KPI Summary Cards
    ws_summary["B12"] = "Class Average:"
    ws_summary["C12"] = f"{data['class_average_percentage']}%"
    ws_summary["C12"].font = Font(bold=True, size=13, color="2B6CB0")

    ws_summary["B13"] = "Participation:"
    ws_summary["C13"] = f"{data['total_participated']} / {data['total_enrolled']} Students ({data['total_absent']} Absent)"

    ws_summary["B14"] = "High Mastery (>=80%):"
    ws_summary["C14"] = f"{data['mastery_counts']['high']} Students"
    ws_summary["C14"].font = Font(color="2E7D32", bold=True)

    ws_summary["B15"] = "Moderate Help (50-79%):"
    ws_summary["C15"] = f"{data['mastery_counts']['moderate']} Students"
    ws_summary["C15"].font = Font(color="E65100", bold=True)

    ws_summary["B16"] = "Critical Support (<50%):"
    ws_summary["C16"] = f"{data['mastery_counts']['low']} Students"
    ws_summary["C16"].font = Font(color="C62828", bold=True)

    for r in range(12, 17):
        ws_summary[f"B{r}"].border = thin_border
        ws_summary[f"C{r}"].border = thin_border

    # -------------------------------------------------------------
    # SHEET 2: Student Scorecard Matrix
    # -------------------------------------------------------------
    ws_scores = wb.create_sheet(title="Student Scorecards")
    ws_scores.views.sheetView[0].showGridLines = True

    # Headers
    q_count = data["question_count"]
    headers = ["Roll No", "Student Name"] + [f"Q{i} (Key: {data['answer_key'].get(str(i), '-')})" for i in range(1, q_count + 1)] + ["Score", "Total", "Percentage", "Mastery Category"]
    ws_scores.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws_scores.cell(row=1, column=col_idx)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_scores.row_dimensions[1].height = 28

    # Rows
    for r_idx, sc in enumerate(data["scorecards"], start=2):
        row_vals = [sc["roll_number"], sc["name"]]
        for q_seq in range(1, q_count + 1):
            b = sc["breakdown"].get(q_seq, {"given": "-"})
            row_vals.append(b["given"])
        row_vals.extend([sc["score"], sc["total_questions"], f"{sc['percentage']}%", sc["mastery"]])
        ws_scores.append(row_vals)

        # Style question cells with green/red highlight
        for col_idx in range(1, len(headers) + 1):
            c = ws_scores.cell(row=r_idx, column=col_idx)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center" if col_idx != 2 else "left", vertical="center")

            # Question response color
            if 3 <= col_idx <= 2 + q_count:
                q_num = col_idx - 2
                is_correct = sc["breakdown"].get(q_num, {}).get("is_correct", False)
                given = sc["breakdown"].get(q_num, {}).get("given", "-")
                if given != "-" and is_correct:
                    c.fill = CORRECT_FILL
                    c.font = Font(color="155724", bold=True)
                elif given != "-":
                    c.fill = WRONG_FILL
                    c.font = Font(color="721C24")

    # -------------------------------------------------------------
    # SHEET 3: Item & Distractor Analysis
    # -------------------------------------------------------------
    ws_items = wb.create_sheet(title="Item Distractor Analysis")
    ws_items.views.sheetView[0].showGridLines = True

    item_headers = ["Question", "Correct Answer", "Accuracy %", "Total Responses", "Option A", "Option B", "Option C", "Option D"]
    ws_items.append(item_headers)

    for col_idx in range(1, len(item_headers) + 1):
        cell = ws_items.cell(row=1, column=col_idx)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for r_idx, q in enumerate(data["questions"], start=2):
        ws_items.append([
            f"Q{q['sequence']}",
            q["correct_answer"] or "N/A",
            f"{q['accuracy_percentage']}%",
            q["total_responses"],
            q["distractors"].get("A", 0),
            q["distractors"].get("B", 0),
            q["distractors"].get("C", 0),
            q["distractors"].get("D", 0)
        ])
        for col_idx in range(1, len(item_headers) + 1):
            c = ws_items.cell(row=r_idx, column=col_idx)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths for all sheets
    for ws in [ws_summary, ws_scores, ws_items]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save workbook
    safe_title = "".join(c for c in data["quiz_name"] if c.isalnum() or c in (' ', '_', '-')).strip()
    file_name = f"Assessment_{safe_title}_{data['class_label']}_{data['session_date']}.xlsx"
    file_path = os.path.join(EXPORT_DIR, file_name)
    wb.save(file_path)
    return file_path


def generate_merged_pdf(session_id: str, custom_answer_key: Optional[Dict[str, str]] = None) -> str:
    """
    Generates a unified PDF dossier:
    - Page 1: Executive Assessment Scorecard & Question Accuracy Table (Branded ReportLab)
    - Page 2+: The original Question Paper PDF uploaded in Step 2 (if available)
    Returns the absolute path to the generated merged PDF.
    """
    data = calculate_session_analysis(session_id, custom_answer_key)
    safe_title = "".join(c for c in data["quiz_name"] if c.isalnum() or c in (' ', '_', '-')).strip()
    scorecard_pdf_path = os.path.join(EXPORT_DIR, f"temp_scorecard_{session_id}.pdf")
    merged_pdf_name = f"Dossier_{safe_title}_{data['class_label']}_{data['session_date']}.pdf"
    merged_pdf_path = os.path.join(EXPORT_DIR, merged_pdf_name)

    # 1. Build Page 1: Scorecard Summary PDF via ReportLab
    doc = SimpleDocTemplate(scorecard_pdf_path, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'MainTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1A365D'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4A5568'),
        spaceAfter=14
    )
    section_title = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2B6CB0'),
        spaceBefore=10,
        spaceAfter=6
    )

    elements = []

    # Header
    elements.append(Paragraph("Pi Jam Foundation — Formative Assessment Dossier", title_style))
    elements.append(Paragraph(f"<b>Quiz:</b> {data['quiz_name']} &nbsp;|&nbsp; <b>Class:</b> {data['class_label']} &nbsp;|&nbsp; <b>School:</b> {data['school_name']} &nbsp;|&nbsp; <b>Date:</b> {data['session_date']}", subtitle_style))

    # KPI Table
    kpi_data = [
        ["Class Average", "Participation Rate", "High Mastery (>=80%)", "Moderate Help (50-79%)", "Critical Support (<50%)"],
        [f"{data['class_average_percentage']}%", f"{data['total_participated']}/{data['total_enrolled']}", f"{data['mastery_counts']['high']}", f"{data['mastery_counts']['moderate']}", f"{data['mastery_counts']['low']}"]
    ]
    kpi_table = Table(kpi_data, colWidths=[108, 108, 108, 108, 108])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8.5),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#EDF2F7')),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 11),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E0'))
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 12))

    # Questions Accuracy & Distractors
    elements.append(Paragraph("Question Item & Distractor Analysis", section_title))
    q_table_data = [["Q#", "Key", "Accuracy %", "Correct / Total", "A", "B", "C", "D"]]
    for q in data["questions"]:
        q_table_data.append([
            f"Q{q['sequence']}",
            q['correct_answer'] or "-",
            f"{q['accuracy_percentage']}%",
            f"{q['correct_count']} / {q['total_responses']}",
            str(q['distractors'].get('A', 0)),
            str(q['distractors'].get('B', 0)),
            str(q['distractors'].get('C', 0)),
            str(q['distractors'].get('D', 0))
        ])

    q_table = Table(q_table_data, colWidths=[40, 45, 80, 115, 65, 65, 65, 65])
    q_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2B6CB0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(q_table)
    elements.append(Spacer(1, 14))

    # Notice if Question Paper is Attached
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT pdf_path, pdf_filename FROM sessions WHERE session_id = ?", (session_id,))
    s_row = cursor.fetchone()
    conn.close()

    uploaded_pdf_path = s_row["pdf_path"] if s_row and s_row["pdf_path"] else None

    if uploaded_pdf_path and os.path.exists(uploaded_pdf_path):
        elements.append(Paragraph(f"<i>* Attached Source Question Paper: <b>{s_row['pdf_filename']}</b> follows immediately on page 2.</i>", subtitle_style))
    else:
        elements.append(Paragraph("<i>* Assessment conducted in Blackboard / Oral mode (no source PDF attached).</i>", subtitle_style))

    # Build the ReportLab summary
    doc.build(elements)

    # 2. Merge with original Question Paper PDF (if uploaded) using pypdf
    merger = PdfWriter()
    merger.append(scorecard_pdf_path)

    if uploaded_pdf_path and os.path.exists(uploaded_pdf_path):
        try:
            merger.append(uploaded_pdf_path)
        except Exception as e:
            print(f"Warning: Could not merge original PDF '{uploaded_pdf_path}': {e}")

    merger.write(merged_pdf_path)
    merger.close()

    # Clean up temp file
    if os.path.exists(scorecard_pdf_path):
        try:
            os.remove(scorecard_pdf_path)
        except Exception:
            pass

    return merged_pdf_path
